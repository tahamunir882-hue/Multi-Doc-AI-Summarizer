import io

SUPPORTED = ["pdf", "docx", "pptx", "txt", "md", "csv", "xlsx", "png", "jpg", "jpeg", "mp3", "wav", "m4a", "mp4"]
# "\f" separates pages / slides so answers can cite page numbers.


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        import fitz  # PyMuPDF
        with fitz.open(stream=data, filetype="pdf") as doc:
            pages = [p.get_text() for p in doc]
        text = "\f".join(pages).strip()
        if len(text.replace("\f", "")) < 50:          # scanned PDF -> OCR
            from src.ocr import ocr_pdf
            text = ocr_pdf(data)
        return text
    if ext == "docx":
        from docx import Document
        return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs).strip()
    if ext == "pptx":
        from pptx import Presentation
        slides = []
        for s in Presentation(io.BytesIO(data)).slides:
            slides.append("\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame))
        return "\f".join(slides).strip()
    if ext in ("csv", "xlsx"):
        import pandas as pd
        if ext == "csv":
            frames = {"data": pd.read_csv(io.BytesIO(data))}
        else:
            frames = pd.read_excel(io.BytesIO(data), sheet_name=None)
        return "\f".join(f"Sheet: {n} ({len(df)} rows)\n{df.head(300).to_string()}"
                         for n, df in frames.items())
    if ext in ("mp3", "wav", "m4a", "mp4"):
        from src.transcribe import transcribe
        return transcribe(filename, data)
    if ext in ("png", "jpg", "jpeg"):
        from src.ocr import ocr_image
        return ocr_image(data)
    if ext in ("txt", "md"):
        for enc in ("utf-8", "utf-16", "cp1256", "latin-1"):
            try:
                return data.decode(enc).strip()
            except UnicodeDecodeError:
                continue
    raise ValueError(f"Unsupported file type: .{ext}")
